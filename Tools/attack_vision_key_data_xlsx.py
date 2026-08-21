#!/usr/bin/env python3
"""Extract attack-vision analysis signals into one time-aligned XLSX sheet."""

import argparse
import bisect
import math
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


TOPIC_FIELDS = {
	"TrajectorySP": ["velocity[0]", "velocity[1]", "velocity[2]"],
	"LocalPosition": ["x", "y", "z", "vx", "vy", "vz", "dist_bottom", "dist_bottom_valid", "z_valid", "v_z_valid"],
	"LocalPositionSP": ["z", "vx", "vy", "vz"],
	"ControlMode": ["flag_armed", "flag_multicopter_position_control_enabled", "flag_control_auto_enabled"],
	"VehicleStatus": ["arming_state", "nav_state", "failsafe"],
	"LandDetected": ["ground_contact", "maybe_landed", "landed", "in_descend"],
	"HoverThrust": ["hover_thrust", "valid"],
	"ThrustSP": ["xyz[0]", "xyz[1]", "xyz[2]"],
	"ActuatorMotors": [f"control[{i}]" for i in range(12)],
	"Battery": ["voltage_v", "current_a", "remaining", "warning", "connected"],
	"EstimatorStatus": ["pos_vert_accuracy", "vel_test_ratio", "hgt_test_ratio", "filter_fault_flags"],
	"FailsafeFlags": ["local_altitude_invalid", "local_velocity_invalid", "battery_warning", "battery_unhealthy"],
}


HEADERS = [
	"时间（time_s）", "分析阶段（analysis_segment）", "阶段编号（segment_index）",
	"模块状态（module_state）", "锁定有效（lock_active）", "图像帧有效（frame_valid）",
	"吊舱LOS有效（los_gimbal_valid）", "NED目标向量有效（target_vec_ned_valid）", "图像帧龄（frame_age_ms）",
	"像素水平偏差（pix_offset_x）", "像素垂直偏差（pix_offset_y）",
	"吊舱横滚角/度（gimbal_roll_deg）", "吊舱俯仰角/度（gimbal_pitch_deg）", "吊舱偏航角/度（gimbal_yaw_deg）",
	"吊舱LOS-X（los_gimbal_x）", "吊舱LOS-Y（los_gimbal_y）", "吊舱LOS-Z（los_gimbal_z）",
	"NED目标向量-N（target_vec_ned_x）", "NED目标向量-E（target_vec_ned_y）", "NED目标向量-D（target_vec_ned_z）",
	"LOS俯角/度-向下为正（los_pitch_deg_down_positive）",
	"目标垂直关系（target_relation:0水平1低2高）", "制导指令有效（guidance_command_valid）",
	"本地高度有效（local_height_valid）", "整形使用本地高度（local_height）",
	"整形前下降速度（vz_raw_ned）", "高度缩放（height_scale）", "角度缩放（angle_scale）",
	"下降联合缩放（descent_scale）", "整形后下降速度（vz_cmd_ned）", "制导计算时间戳（guidance_timestamp_us）",
	"期望北向速度（trajectory_vx）", "期望东向速度（trajectory_vy）", "期望下降速度（trajectory_vz）",
	"期望水平速度（trajectory_vxy）", "期望三维速度（trajectory_vxyz）",
	"实际北向速度（local_vx）", "实际东向速度（local_vy）", "实际下降速度（local_vz）",
	"实际水平速度（local_vxy）", "垂向速度误差（vz_actual_minus_cmd）",
	"实际相对高度（altitude=-local_z）", "期望相对高度（altitude_sp=-local_position_sp_z）", "高度误差（altitude_actual_minus_sp）",
	"本地位置Z（local_z）", "本地位置期望Z（local_position_sp_z）",
	"对地距离-仅参考（dist_bottom）", "对地距离有效（dist_bottom_valid）", "高度有效（z_valid）", "垂速有效（v_z_valid）",
	"悬停推力估计（hover_thrust）", "悬停推力有效（hover_thrust_valid）",
	"推力期望X（thrust_sp_x）", "推力期望Y（thrust_sp_y）", "推力期望Z（thrust_sp_z）", "推力期望幅值（thrust_sp_norm）",
]

HEADERS += [f"电机{i + 1}输出（actuator_motor_{i}）" for i in range(12)]
HEADERS += [
	"最大电机输出（motor_max）", "最小电机输出（motor_min）", "电机输出跨度（motor_range）",
	"电池电压（battery_voltage_v）", "电池电流（battery_current_a）", "电池剩余比例（battery_remaining）", "电池告警（battery_warning）",
	"已解锁（flag_armed）", "位置控制使能（position_control_enabled）", "自动控制使能（auto_control_enabled）",
	"解锁状态（arming_state）", "导航状态（nav_state）", "整机失效保护（vehicle_failsafe）",
	"着地（landed）", "可能着地（maybe_landed）", "地面接触（ground_contact）", "正在下降（in_descend）",
	"垂直位置精度（pos_vert_accuracy）", "速度检验比（vel_test_ratio）", "高度检验比（hgt_test_ratio）", "估计器故障标志（filter_fault_flags）",
	"本地高度无效（local_altitude_invalid）", "本地速度无效（local_velocity_invalid）", "失效保护电池告警（failsafe_battery_warning）", "电池不健康（battery_unhealthy）",
]


def read_topic(workbook, sheet_name, fields):
	if sheet_name not in workbook.sheetnames:
		return {"times": [], "rows": [], "fields": fields}
	sheet = workbook[sheet_name]
	headers = list(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
	indices = {field: headers.index(field) for field in fields if field in headers}
	time_index = headers.index("time_s")
	times, rows = [], []
	for row in sheet.iter_rows(min_row=2, values_only=True):
		if row[time_index] is None:
			continue
		times.append(float(row[time_index]))
		rows.append({field: row[index] for field, index in indices.items()})
	return {"times": times, "rows": rows, "fields": fields}


def previous_sample(topic, time_s):
	index = bisect.bisect_right(topic["times"], time_s) - 1
	return topic["rows"][index] if index >= 0 else {}


def finite(value):
	return isinstance(value, (int, float)) and math.isfinite(value)


def norm(*values):
	return math.sqrt(sum(value * value for value in values)) if all(finite(value) for value in values) else None


def segment_at(segments, time_s):
	for segment_type, index, start_s, end_s in segments:
		if start_s <= time_s <= end_s:
			return segment_type, index
	return "OTHER", None


def build_row(attack, samples, segment):
	trajectory = samples["TrajectorySP"]
	local = samples["LocalPosition"]
	local_sp = samples["LocalPositionSP"]
	hover = samples["HoverThrust"]
	thrust = samples["ThrustSP"]
	motors = [samples["ActuatorMotors"].get(f"control[{i}]") for i in range(12)]
	valid_motors = [value for value in motors if finite(value) and value >= 0]
	control = samples["ControlMode"]
	status = samples["VehicleStatus"]
	land = samples["LandDetected"]
	battery = samples["Battery"]
	estimator = samples["EstimatorStatus"]
	failsafe = samples["FailsafeFlags"]

	los_n = attack.get("target_vec_ned_x")
	los_e = attack.get("target_vec_ned_y")
	los_d = attack.get("target_vec_ned_z")
	los_pitch = math.degrees(math.atan2(los_d, math.hypot(los_n, los_e))) if all(finite(v) for v in (los_n, los_e, los_d)) else None
	cmd_vx, cmd_vy, cmd_vz = (trajectory.get(f"velocity[{i}]") for i in range(3))
	vx, vy, vz = local.get("vx"), local.get("vy"), local.get("vz")
	z, z_sp = local.get("z"), local_sp.get("z")
	altitude = -z if finite(z) else None
	altitude_sp = -z_sp if finite(z_sp) else None
	thrust_values = [thrust.get(f"xyz[{i}]") for i in range(3)]

	row = [
		attack["time_s"], segment[0], segment[1], attack.get("module_state"), attack.get("lock_active"), attack.get("frame_valid"),
		attack.get("los_gimbal_valid"), attack.get("target_vec_ned_valid"), attack.get("frame_age_ms"),
		attack.get("pix_offset_x"), attack.get("pix_offset_y"),
		math.degrees(attack["gimbal_roll_rad"]) if finite(attack.get("gimbal_roll_rad")) else None,
		math.degrees(attack["gimbal_pitch_rad"]) if finite(attack.get("gimbal_pitch_rad")) else None,
		math.degrees(attack["gimbal_yaw_rad"]) if finite(attack.get("gimbal_yaw_rad")) else None,
		attack.get("los_gimbal_x"), attack.get("los_gimbal_y"), attack.get("los_gimbal_z"), los_n, los_e, los_d, los_pitch,
		attack.get("target_relation"), attack.get("guidance_command_valid"), attack.get("local_height_valid"),
		attack.get("local_height"), attack.get("vz_raw_ned"), attack.get("height_scale"), attack.get("angle_scale"),
		attack.get("descent_scale"), attack.get("vz_cmd_ned"), attack.get("guidance_timestamp"),
		cmd_vx, cmd_vy, cmd_vz, norm(cmd_vx, cmd_vy), norm(cmd_vx, cmd_vy, cmd_vz),
		vx, vy, vz, norm(vx, vy), vz - cmd_vz if finite(vz) and finite(cmd_vz) else None,
		altitude, altitude_sp, altitude - altitude_sp if finite(altitude) and finite(altitude_sp) else None,
		z, z_sp, local.get("dist_bottom"), local.get("dist_bottom_valid"), local.get("z_valid"), local.get("v_z_valid"),
		hover.get("hover_thrust"), hover.get("valid"), *thrust_values, norm(*thrust_values), *motors,
		max(valid_motors) if valid_motors else None, min(valid_motors) if valid_motors else None,
		max(valid_motors) - min(valid_motors) if valid_motors else None,
		battery.get("voltage_v"), battery.get("current_a"), battery.get("remaining"), battery.get("warning"),
		control.get("flag_armed"), control.get("flag_multicopter_position_control_enabled"), control.get("flag_control_auto_enabled"),
		status.get("arming_state"), status.get("nav_state"), status.get("failsafe"),
		land.get("landed"), land.get("maybe_landed"), land.get("ground_contact"), land.get("in_descend"),
		estimator.get("pos_vert_accuracy"), estimator.get("vel_test_ratio"), estimator.get("hgt_test_ratio"), estimator.get("filter_fault_flags"),
		failsafe.get("local_altitude_invalid"), failsafe.get("local_velocity_invalid"), failsafe.get("battery_warning"), failsafe.get("battery_unhealthy"),
	]
	return row


def export(source_path, output_path):
	source = load_workbook(source_path, read_only=True, data_only=True)
	attack_topic = read_topic(source, "AttackVision", [
		"time_s", "gimbal_roll_rad", "gimbal_pitch_rad", "gimbal_yaw_rad", "los_gimbal_x", "los_gimbal_y", "los_gimbal_z",
		"target_vec_ned_x", "target_vec_ned_y", "target_vec_ned_z", "frame_age_ms", "pix_offset_x", "pix_offset_y",
		"lock_active", "frame_valid", "los_gimbal_valid", "target_vec_ned_valid", "module_state",
		"guidance_command_valid", "local_height_valid", "target_relation", "local_height", "vz_raw_ned",
		"height_scale", "angle_scale", "descent_scale", "vz_cmd_ned", "guidance_timestamp",
	])
	topics = {name: read_topic(source, name, fields) for name, fields in TOPIC_FIELDS.items()}
	segment_sheet = source["Analysis_Segments"]
	segment_headers = list(next(segment_sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
	segment_columns = [segment_headers.index(name) for name in ("segment_type", "segment_index", "start_s", "end_s")]
	segments = [tuple(row[index] for index in segment_columns) for row in segment_sheet.iter_rows(min_row=2, values_only=True)]

	workbook = Workbook(write_only=False)
	sheet = workbook.active
	sheet.title = "关键分析数据_KeyData"
	sheet.sheet_view.showGridLines = False
	sheet.append(HEADERS)
	for time_s, attack in zip(attack_topic["times"], attack_topic["rows"]):
		attack["time_s"] = time_s
		samples = {name: previous_sample(topic, time_s) for name, topic in topics.items()}
		sheet.append(build_row(attack, samples, segment_at(segments, time_s)))

	last_row, last_col = sheet.max_row, sheet.max_column
	header_fill = PatternFill("solid", fgColor="1F4E78")
	header_font = Font(name="Microsoft YaHei", color="FFFFFF", bold=True, size=10)
	for cell in sheet[1]:
		cell.fill = header_fill
		cell.font = header_font
		cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
		cell.border = Border(bottom=Side(style="medium", color="9EADBA"))
	sheet.row_dimensions[1].height = 48
	sheet.freeze_panes = "D2"
	sheet.auto_filter.ref = f"A1:{get_column_letter(last_col)}{last_row}"
	sheet.sheet_properties.pageSetUpPr.fitToPage = True
	sheet.sheet_view.zoomScale = 75

	for column in range(1, last_col + 1):
		letter = get_column_letter(column)
		sheet.column_dimensions[letter].width = 18 if column > 3 else (13 if column == 1 else 16)
		for cell in sheet.iter_cols(min_col=column, max_col=column, min_row=2, max_row=last_row):
			for item in cell:
				item.font = Font(name="Microsoft YaHei", size=9)
				item.number_format = "0.000"

	attack_fill = PatternFill("solid", fgColor="E2F0D9")
	segment_col = HEADERS.index("分析阶段（analysis_segment）") + 1
	segment_letter = get_column_letter(segment_col)
	sheet.conditional_formatting.add(
		f"A2:{get_column_letter(last_col)}{last_row}",
		FormulaRule(formula=[f'${segment_letter}2="ATTACK"'], fill=attack_fill),
	)

	table = Table(displayName="AttackVisionKeyData", ref=f"A1:{get_column_letter(last_col)}{last_row}")
	table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
	sheet.add_table(table)
	output_path.parent.mkdir(parents=True, exist_ok=True)
	workbook.save(output_path)
	source.close()
	workbook.close()

	verified = load_workbook(output_path, read_only=True, data_only=True)
	verified_sheet = verified["关键分析数据_KeyData"]
	if verified_sheet.max_row != len(attack_topic["times"]) + 1 or verified_sheet.max_column != len(HEADERS):
		raise RuntimeError("Output workbook verification failed")
	verified.close()


def main():
	parser = argparse.ArgumentParser(description=__doc__)
	parser.add_argument("source", type=Path, help="Full XLSX exported by attack_vision_ulog_to_xlsx.py")
	parser.add_argument("-o", "--output", type=Path, help="Output path")
	args = parser.parse_args()
	output = args.output or Path("outputs") / f"{args.source.stem}_关键分析数据.xlsx"
	export(args.source, output)
	print(output.resolve())


if __name__ == "__main__":
	main()
