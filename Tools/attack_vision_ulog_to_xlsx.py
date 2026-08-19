#!/usr/bin/env python3
"""Export attack_vision flight data from PX4 ULog files to Excel workbooks."""

import argparse
import math
import sys
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pyulog import ULog


TOPICS = (
	("attack_vision_status", "AttackVision"),
	("trajectory_setpoint", "TrajectorySP"),
	("vehicle_local_position", "LocalPosition"),
	("vehicle_local_position_setpoint", "LocalPositionSP"),
	("vehicle_status", "VehicleStatus"),
	("vehicle_control_mode", "ControlMode"),
	("vehicle_attitude", "Attitude"),
	("manual_control_setpoint", "ManualControl"),
	("hover_thrust_estimate", "HoverThrust"),
	("vehicle_thrust_setpoint", "ThrustSP"),
	("actuator_motors", "ActuatorMotors"),
	("battery_status", "Battery"),
	("vehicle_gps_position", "GPS"),
	("estimator_status", "EstimatorStatus"),
	("estimator_aid_src_baro_hgt", "BaroHeightAid"),
	("failsafe_flags", "FailsafeFlags"),
	("vehicle_land_detected", "LandDetected"),
)

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
TITLE_FILL = PatternFill("solid", fgColor="DDEBF7")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def excel_value(value):
	"""Convert numpy and non-finite values into values accepted by Excel."""
	if isinstance(value, np.generic):
		value = value.item()
	if isinstance(value, float):
		return value if math.isfinite(value) else None
	if isinstance(value, bytes):
		return value.decode("utf-8", errors="replace")
	return value


def dataset(ulog, topic, multi_id=0):
	for item in ulog.data_list:
		if item.name == topic and item.multi_id == multi_id:
			return item.data
	return None


def previous(data, field, timestamps, default=np.nan):
	if data is None or field not in data or not len(data["timestamp"]):
		return np.full(len(timestamps), default)

	source_time = np.asarray(data["timestamp"], dtype=np.int64)
	indices = np.searchsorted(source_time, timestamps, side="right") - 1
	valid = indices >= 0
	result = np.full(len(timestamps), default, dtype=float)
	result[valid] = np.asarray(data[field])[indices[valid]]
	return result


def contiguous_segments(timestamps, mask, min_duration_s, max_gap_s):
	indices = np.flatnonzero(mask)
	if not len(indices):
		return []

	result = []
	start = previous_index = indices[0]
	for index in indices[1:]:
		if index != previous_index + 1 or timestamps[index] - timestamps[previous_index] > max_gap_s * 1e6:
			if timestamps[previous_index] - timestamps[start] >= min_duration_s * 1e6:
				result.append((start, previous_index))
			start = index
		previous_index = index

	if timestamps[previous_index] - timestamps[start] >= min_duration_s * 1e6:
		result.append((start, previous_index))
	return result


def finite_mean(values):
	values = np.asarray(values, dtype=float)
	values = values[np.isfinite(values)]
	return float(np.mean(values)) if len(values) else None


def finite_min(values):
	values = np.asarray(values, dtype=float)
	values = values[np.isfinite(values)]
	return float(np.min(values)) if len(values) else None


def finite_max(values):
	values = np.asarray(values, dtype=float)
	values = values[np.isfinite(values)]
	return float(np.max(values)) if len(values) else None


def style_header(worksheet):
	for cell in worksheet[1]:
		cell.fill = HEADER_FILL
		cell.font = HEADER_FONT
		cell.alignment = Alignment(horizontal="center", vertical="center")
	worksheet.row_dimensions[1].height = 24


def configure_data_sheet(worksheet, column_count, data_row_count):
	worksheet.freeze_panes = "B2"
	worksheet.auto_filter.ref = f"A1:{get_column_letter(column_count)}{data_row_count + 1}"
	worksheet.sheet_view.showGridLines = False
	worksheet.column_dimensions["A"].width = 14
	for column in range(2, column_count + 1):
		worksheet.column_dimensions[get_column_letter(column)].width = 18
	style_header(worksheet)
	for cell in worksheet[1]:
		cell.alignment = Alignment(text_rotation=45, horizontal="center", vertical="bottom")
	worksheet.row_dimensions[1].height = 92


def write_topic_sheets(workbook, ulog, topic, sheet_name):
	written = []
	for item in (item for item in ulog.data_list if item.name == topic):
		name = (sheet_name + (f"_{item.multi_id}" if item.multi_id else ""))[:31]
		worksheet = workbook.create_sheet(name)
		fields = list(item.data.keys())
		timestamps = np.asarray(item.data["timestamp"], dtype=np.int64)
		arrays = [np.asarray(item.data[field]) for field in fields]
		worksheet.append(["time_s", *fields])

		for row_index, timestamp in enumerate(timestamps):
			worksheet.append([
				(int(timestamp) - ulog.start_timestamp) * 1e-6,
				*(excel_value(array[row_index]) for array in arrays),
			])

		configure_data_sheet(worksheet, len(fields) + 1, len(timestamps))
		for cell in worksheet["A"][1:]:
			cell.number_format = "0.000"
		written.append((name, topic, len(timestamps)))
	return written


def build_segment_rows(ulog):
	rows = []
	attack = dataset(ulog, "attack_vision_status")
	local_position = dataset(ulog, "vehicle_local_position")
	local_setpoint = dataset(ulog, "vehicle_local_position_setpoint")
	trajectory = dataset(ulog, "trajectory_setpoint")
	status = dataset(ulog, "vehicle_status")
	land = dataset(ulog, "vehicle_land_detected")

	if attack is not None and local_position is not None:
		timestamps = np.asarray(attack["timestamp"], dtype=np.int64)
		active = ((np.asarray(attack["module_state"]) == 2)
			  & (np.asarray(attack["lock_active"]) > 0)
			  & (np.asarray(attack["frame_valid"]) > 0)
			  & (np.asarray(attack["target_vec_ned_valid"]) > 0))

		for number, (start, end) in enumerate(contiguous_segments(timestamps, active, 1.0, 0.2), 1):
			times = timestamps[start:end + 1]
			altitude = -previous(local_position, "z", times)
			actual_vz = previous(local_position, "vz", times)
			command_vz = previous(trajectory, "velocity[2]", times)
			target_z = np.asarray(attack["target_vec_ned_z"], dtype=float)[start:end + 1]
			los_pitch = np.degrees(np.arcsin(np.clip(target_z, -1.0, 1.0)))
			rows.append(["ATTACK", number, times[0], times[-1], altitude[0], altitude[-1],
				finite_min(altitude), finite_mean(los_pitch), finite_mean(command_vz),
				finite_max(command_vz), finite_mean(actual_vz), finite_max(actual_vz),
				finite_mean(actual_vz - command_vz)])

	if local_position is not None and local_setpoint is not None and status is not None:
		timestamps = np.asarray(local_position["timestamp"], dtype=np.int64)
		nav_state = previous(status, "nav_state", timestamps)
		armed = previous(status, "arming_state", timestamps) == 2
		airborne = previous(land, "landed", timestamps, 0) < 0.5
		setpoint_vz = previous(local_setpoint, "vz", timestamps)
		setpoint_z = previous(local_setpoint, "z", timestamps)
		hold = (armed & airborne & (nav_state == 2) & np.isfinite(setpoint_z)
			& np.isfinite(setpoint_vz) & (np.abs(setpoint_vz) < 0.08))
		altitude = -np.asarray(local_position["z"], dtype=float)
		actual_vz = np.asarray(local_position["vz"], dtype=float)

		for number, (start, end) in enumerate(contiguous_segments(timestamps, hold, 3.0, 0.5), 1):
			slice_ = slice(start, end + 1)
			rows.append(["POSITION_HOLD", number, timestamps[start], timestamps[end],
				altitude[start], altitude[end], finite_min(altitude[slice_]), None,
				finite_mean(setpoint_vz[slice_]), finite_max(setpoint_vz[slice_]),
				finite_mean(actual_vz[slice_]), finite_max(actual_vz[slice_]),
				finite_mean(actual_vz[slice_] - setpoint_vz[slice_])])

	for row in rows:
		start_timestamp, end_timestamp = row[2], row[3]
		row[2:4] = [
			(start_timestamp - ulog.start_timestamp) * 1e-6,
			(end_timestamp - ulog.start_timestamp) * 1e-6,
			(end_timestamp - start_timestamp) * 1e-6,
		]
	return rows


def write_overview(workbook, ulog, source, topic_rows, segment_rows):
	worksheet = workbook.active
	worksheet.title = "Overview"
	worksheet.sheet_view.showGridLines = False
	worksheet.merge_cells("A1:D1")
	worksheet["A1"] = "PX4 Flight Log Analysis Workbook"
	worksheet["A1"].font = Font(size=15, bold=True, color="17365D")
	worksheet["A1"].fill = TITLE_FILL
	worksheet.append([])
	worksheet.append(["Log information", "Value"])
	for label, value in (
		("Source ULog", source.name),
		("File size (MB)", round(source.stat().st_size / 1024 / 1024, 3)),
		("Log duration (s)", round((ulog.last_timestamp - ulog.start_timestamp) * 1e-6, 3)),
		("Start timestamp (us)", int(ulog.start_timestamp)),
		("Software version", ulog.msg_info_dict.get("ver_sw", "")),
		("Hardware", ulog.msg_info_dict.get("ver_hw", "")),
		("Time convention", "time_s is seconds from ULog start; NED vz > 0 means descending"),
	):
		worksheet.append([label, value])

	worksheet.append([])
	topic_header_row = worksheet.max_row + 1
	worksheet.append(["Exported sheet", "Source topic", "Rows"])
	for row in topic_rows:
		worksheet.append(row)
	worksheet.append([])
	segment_header_row = worksheet.max_row + 1
	worksheet.append(["Detected segment", "Count"])
	worksheet.append(["Attack Vision", sum(row[0] == "ATTACK" for row in segment_rows)])
	worksheet.append(["Position hold", sum(row[0] == "POSITION_HOLD" for row in segment_rows)])
	for row_number in (3, topic_header_row, segment_header_row):
		for cell in worksheet[row_number]:
			cell.fill = HEADER_FILL
			cell.font = HEADER_FONT
	worksheet.column_dimensions["A"].width = 28
	worksheet.column_dimensions["B"].width = 66
	worksheet.column_dimensions["C"].width = 14
	worksheet.freeze_panes = "A3"


def write_parameters(workbook, ulog):
	worksheet = workbook.create_sheet("Parameters")
	worksheet.append(["Parameter", "Initial value", "Group"])
	for name, value in sorted(ulog.initial_parameters.items()):
		if name.startswith(("AV_", "AAATTK")):
			group = "Attack Vision"
		elif name.startswith("MPC_"):
			group = "Position control"
		elif name.startswith(("EKF2_", "SENS_")):
			group = "Estimator"
		else:
			group = "Other"
		worksheet.append([name, excel_value(value), group])
	style_header(worksheet)
	worksheet.freeze_panes = "A2"
	worksheet.auto_filter.ref = f"A1:C{worksheet.max_row}"
	worksheet.sheet_view.showGridLines = False
	worksheet.column_dimensions["A"].width = 28
	worksheet.column_dimensions["B"].width = 20
	worksheet.column_dimensions["C"].width = 22


def write_segments(workbook, rows):
	headers = ("segment_type", "segment_index", "start_s", "end_s", "duration_s",
		"alt_start_m", "alt_end_m", "alt_min_m", "los_pitch_mean_deg",
		"vz_cmd_mean_mps", "vz_cmd_max_mps", "vz_actual_mean_mps",
		"vz_actual_max_mps", "vz_error_mean_mps")
	worksheet = workbook.create_sheet("Analysis_Segments")
	worksheet.append(headers)
	for row in rows:
		worksheet.append([excel_value(value) for value in row])
	configure_data_sheet(worksheet, len(headers), len(rows))
	worksheet.freeze_panes = "A2"
	for cell in worksheet[1]:
		cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
	worksheet.row_dimensions[1].height = 38


def verify_workbook(path):
	workbook = load_workbook(path, read_only=True, data_only=False)
	required = {"Overview", "Parameters", "Analysis_Segments", "AttackVision", "LocalPosition"}
	missing = required - set(workbook.sheetnames)
	if missing:
		raise RuntimeError(f"missing worksheets: {', '.join(sorted(missing))}")
	for worksheet in workbook.worksheets:
		if worksheet.max_row < 1 or worksheet.max_column < 1:
			raise RuntimeError(f"empty worksheet: {worksheet.title}")
	workbook.close()


def export_log(source, output_dir):
	# Ignore malformed UTF-8 info strings while preserving valid ULog data records.
	ulog = ULog(str(source), None, True)
	workbook = Workbook()
	segment_rows = build_segment_rows(ulog)
	topic_rows = []
	for topic, sheet_name in TOPICS:
		for item in ulog.data_list:
			if item.name == topic:
				name = (sheet_name + (f"_{item.multi_id}" if item.multi_id else ""))[:31]
				topic_rows.append((name, topic, len(item.data["timestamp"])))

	write_overview(workbook, ulog, source, topic_rows, segment_rows)
	write_parameters(workbook, ulog)
	write_segments(workbook, segment_rows)
	for topic, sheet_name in TOPICS:
		write_topic_sheets(workbook, ulog, topic, sheet_name)

	output = output_dir / f"{source.stem}.xlsx"
	workbook.save(output)
	verify_workbook(output)
	return output


def main():
	parser = argparse.ArgumentParser(
		description="Export attack_vision related PX4 ULog data to one Excel workbook per log.")
	parser.add_argument("logs", nargs="+", type=Path, help="input .ulg file(s)")
	args = parser.parse_args()

	repository_root = Path(__file__).resolve().parent.parent
	output_dir = repository_root / "outputs"
	output_dir.mkdir(parents=True, exist_ok=True)
	failures = 0

	for source in args.logs:
		if not source.is_file() or source.suffix.lower() != ".ulg":
			print(f"error: not a ULog file: {source}", file=sys.stderr)
			failures += 1
			continue
		try:
			output = export_log(source.resolve(), output_dir)
			print(output)
		except Exception as error:  # Keep batch exports running after one bad log.
			print(f"error: {source}: {error}", file=sys.stderr)
			failures += 1

	return 1 if failures else 0


if __name__ == "__main__":
	raise SystemExit(main())
